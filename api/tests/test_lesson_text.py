from __future__ import annotations

import io

import fitz
import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from app.services.lesson_text import (
    LessonTextError,
    TextSection,
    chunk_sections,
    extract_sections,
    objective_chunk_text,
)


def _pptx_bytes() -> bytes:
    presentation = Presentation()
    layout = presentation.slide_layouts[1]
    first = presentation.slides.add_slide(layout)
    first.shapes.title.text = "OODA Cycle"
    first.placeholders[1].text = "Observe, Orient, Decide, Act"
    first.notes_slide.notes_text_frame.text = "Stress that orientation drives the loop."

    second = presentation.slides.add_slide(presentation.slide_layouts[5])
    second.shapes.title.text = "Decision-making variables"
    table = second.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    table.cell(0, 0).text = "Uncertainty"
    table.cell(0, 1).text = "Time"
    table.cell(1, 0).text = "Experience"
    table.cell(1, 1).text = "Information"

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Interior Guard", level=1)
    document.add_paragraph("The interior guard protects property and enforces regulations.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "JOOD"
    table.cell(0, 1).text = "Junior Officer of the Day"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pdf_bytes() -> bytes:
    document = fitz.open()
    for text in ("Hazing is prohibited.", "Fraternization policy."):
        page = document.new_page()
        page.insert_text((72, 72), text)
    data = document.tobytes()
    document.close()
    return data


def test_pptx_reads_slides_tables_and_speaker_notes() -> None:
    sections = extract_sections("OODA.pptx", _pptx_bytes())

    assert [section.locator for section in sections] == [{"slide": 1}, {"slide": 2}]
    assert "Observe, Orient, Decide, Act" in sections[0].text
    assert "Speaker notes: Stress that orientation drives the loop." in sections[0].text
    assert "Uncertainty | Time" in sections[1].text


def test_docx_reads_paragraphs_and_tables() -> None:
    sections = extract_sections("guard.docx", _docx_bytes())

    assert sections[0].locator == {"part": 1}
    assert "protects property" in sections[0].text
    assert "JOOD | Junior Officer of the Day" in sections[0].text


def test_pdf_reads_one_section_per_page() -> None:
    sections = extract_sections("conduct.pdf", _pdf_bytes())

    assert [section.locator for section in sections] == [{"page": 1}, {"page": 2}]
    assert "Hazing" in sections[0].text


def test_txt_is_read_as_is() -> None:
    sections = extract_sections("notes.txt", "Four domains of Marine Fitness".encode())
    assert sections[0].text == "Four domains of Marine Fitness"


@pytest.mark.parametrize("filename", ["old.ppt", "old.doc", "roster.xlsx", "photo.png"])
def test_legacy_and_binary_formats_are_skipped(filename: str) -> None:
    with pytest.raises(LessonTextError):
        extract_sections(filename, b"\xd0\xcf\x11\xe0")


def test_corrupt_file_raises_readable_error() -> None:
    with pytest.raises(LessonTextError, match="Could not read"):
        extract_sections("broken.pptx", b"PK\x03\x04garbage")


def test_chunks_follow_slide_boundaries_with_headers() -> None:
    sections = extract_sections("OODA.pptx", _pptx_bytes())
    chunks = chunk_sections(sections, header="Decision Making · OODA.pptx")

    assert [chunk.locator for chunk in chunks] == [{"slide": 1}, {"slide": 2}]
    assert chunks[0].text.startswith("Decision Making · OODA.pptx · slide 1\n")
    assert [chunk.sequence_index for chunk in chunks] == [0, 1]


def test_long_sections_split_with_overlap() -> None:
    body = " ".join(f"Sentence {index} about casualty evacuation." for index in range(400))
    chunks = chunk_sections([TextSection(body, {"page": 3})], header="TCCC · CASEVAC.pdf", max_chars=1000, overlap=100)

    assert len(chunks) > 1
    assert all(chunk.locator == {"page": 3} for chunk in chunks)
    assert all(len(chunk.text) <= 1000 + len("TCCC · CASEVAC.pdf · p.3\n") for chunk in chunks)


def test_objective_chunk_text_carries_class_code_and_async() -> None:
    text = objective_chunk_text(
        "Decision Making",
        {"position": 3, "text": "Define analytical decision-making without omission.", "code": "MCS-CMBH-1002c", "is_async": True},
    )
    assert text == (
        "Decision Making — ELO 3: Define analytical decision-making without omission. "
        "(MCS-CMBH-1002c) [Async]"
    )
