"""Text extraction and chunking for Academy ELO lesson files and ELOs.

Output feeds lesson_chunks, which Discussions searches alongside ndr_segments
and wiki_entries. Nothing here touches the Knowledge pipeline.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Any

from app.services.lesson_files import file_extension, skip_reason

MAX_CHUNK_CHARS = 3000
CHUNK_OVERLAP_CHARS = 200
DOCX_PART_CHARS = 2500


class LessonTextError(ValueError):
    """Raised when a lesson file cannot be read."""


@dataclass(frozen=True)
class TextSection:
    text: str
    locator: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class LessonChunk:
    text: str
    locator: dict[str, Any]
    sequence_index: int


def extract_sections(filename: str, content: bytes) -> list[TextSection]:
    """One section per slide or page (DOCX: per ~2,500 characters of body text)."""
    reason = skip_reason(filename)
    if reason:
        raise LessonTextError(reason)
    extension = file_extension(filename)
    try:
        if extension == ".pdf":
            sections = _pdf_sections(content)
        elif extension == ".pptx":
            sections = _pptx_sections(content)
        elif extension == ".docx":
            sections = _docx_sections(content)
        else:
            sections = _text_sections(content)
    except LessonTextError:
        raise
    except Exception as exc:  # noqa: BLE001 - parser errors vary by library
        raise LessonTextError(f"Could not read {filename}: {exc}") from exc
    return [section for section in sections if section.text.strip()]


def _clean(text: str) -> str:
    lines = [" ".join(line.split()) for line in text.replace("\r", "\n").split("\n")]
    collapsed: list[str] = []
    for line in lines:
        if line or (collapsed and collapsed[-1]):
            collapsed.append(line)
    return "\n".join(collapsed).strip()


def _pdf_sections(content: bytes) -> list[TextSection]:
    import fitz

    with fitz.open(stream=content, filetype="pdf") as document:
        return [
            TextSection(_clean(page.get_text("text")), {"page": index})
            for index, page in enumerate(document, start=1)
        ]


def _shape_texts(shapes: Any) -> list[str]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    texts: list[str] = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            texts.extend(_shape_texts(shape.shapes))
            continue
        if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
            texts.append(shape.text_frame.text)
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    texts.append(" | ".join(cells))
    return texts


def _pptx_sections(content: bytes) -> list[TextSection]:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(content))
    sections: list[TextSection] = []
    for index, slide in enumerate(presentation.slides, start=1):
        body = _clean("\n".join(_shape_texts(slide.shapes)))
        notes = ""
        if slide.has_notes_slide:
            notes = _clean(slide.notes_slide.notes_text_frame.text or "")
        parts = [body] if body else []
        if notes:
            parts.append(f"Speaker notes: {notes}")
        sections.append(TextSection("\n".join(parts), {"slide": index}))
    return sections


def _docx_sections(content: bytes) -> list[TextSection]:
    from docx import Document

    document = Document(io.BytesIO(content))
    blocks: list[str] = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))

    sections: list[TextSection] = []
    current: list[str] = []
    size = 0
    for block in blocks:
        if current and size + len(block) > DOCX_PART_CHARS:
            sections.append(TextSection(_clean("\n".join(current)), {"part": len(sections) + 1}))
            current, size = [], 0
        current.append(block)
        size += len(block) + 1
    if current:
        sections.append(TextSection(_clean("\n".join(current)), {"part": len(sections) + 1}))
    return sections


def _text_sections(content: bytes) -> list[TextSection]:
    return [TextSection(_clean(content.decode("utf-8", errors="replace")), {"part": 1})]


def locator_label(locator: dict[str, Any]) -> str | None:
    if locator.get("slide"):
        return f"slide {locator['slide']}"
    if locator.get("page"):
        return f"p.{locator['page']}"
    if locator.get("part"):
        return f"part {locator['part']}"
    return None


def _split(text: str, max_chars: int, overlap: int) -> list[str]:
    if len(text) <= max_chars:
        return [text]
    pieces: list[str] = []
    start = 0
    while start < len(text):
        end = min(start + max_chars, len(text))
        if end < len(text):
            window = text[start:end]
            cut = max(window.rfind("\n"), window.rfind(". "))
            if cut > max_chars // 2:
                end = start + cut + 1
        pieces.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return [piece for piece in pieces if piece]


def chunk_sections(
    sections: list[TextSection],
    *,
    header: str,
    max_chars: int = MAX_CHUNK_CHARS,
    overlap: int = CHUNK_OVERLAP_CHARS,
) -> list[LessonChunk]:
    """Split sections into chunks that each start with a class/file/slide header."""
    chunks: list[LessonChunk] = []
    for section in sections:
        label = locator_label(section.locator)
        prefix = f"{header} · {label}" if label else header
        for piece in _split(section.text, max_chars, overlap):
            chunks.append(
                LessonChunk(
                    text=f"{prefix}\n{piece}",
                    locator=dict(section.locator),
                    sequence_index=len(chunks),
                )
            )
    return chunks


def objective_chunk_text(lesson_title: str, objective: dict[str, Any]) -> str:
    text = f"{lesson_title} — ELO {objective['position']}: {objective['text']}"
    if objective.get("code"):
        text += f" ({objective['code']})"
    if objective.get("is_async"):
        text += " [Async]"
    return text
